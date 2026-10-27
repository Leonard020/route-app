from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

out = Path(r"C:\dev\route-app\docs\기능정의서_완성.xlsx")
out.parent.mkdir(parents=True, exist_ok=True)

wb = Workbook()
header_fill = PatternFill("solid", fgColor="1F4E79")
header_font = Font(color="FFFFFF", bold=True)
wrap = Alignment(wrap_text=True, vertical="top")
thin = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)


def style_header(ws, cols):
    for col, name in enumerate(cols, 1):
        cell = ws.cell(1, col, name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{chr(64+len(cols))}1"


def add_rows(ws, rows):
    for r, row in enumerate(rows, 2):
        for c, val in enumerate(row, 1):
            cell = ws.cell(r, c, val)
            cell.alignment = wrap
            cell.border = thin


# --- 사용자 기능 ---
ws = wb.active
ws.title = "사용자기능"
style_header(ws, ["No", "1Depth", "2Depth", "Page ID", "Page Name", "Target", "Content", "Description", "백엔드 API", "실시간 연동"])
user_rows = [
    ["1", "공통", "권한 설정", "CM_01_01", "위치 권한 안내", "App", "위치정보 접근 권한 안내", "최초 실행 시 위치 활용 목적 안내 후 OS 권한 요청", "클라이언트 OS", "-"],
    ["2", "공통", "현재 위치 조회", "CM_01_01", "현재 위치 조회", "App", "현재 위치 좌표 및 주소", "GPS 좌표를 주소/인근 장소명으로 변환", "GET /api/v1/places/search", "-"],
    ["3", "공통", "알림 권한", "CM_01_01", "이동 안내 알림 권한", "App", "하차·환승 알림", "길안내 시작 시 알림 권한 확인", "GET/PATCH /api/v1/settings", "-"],
    ["4", "공통", "예외 처리", "ER_01_01", "위치 조회 실패", "Web/App", "권한 거부/측위 실패", "기존 수동 출발지 유지, 장소 검색 이동", "-", "-"],
    ["5", "공통", "예외 처리", "ER_01_01", "실시간 데이터 미수신", "Web/App", "버스·지하철 실시간 미제공", "마지막 정상 응답 + 조회시각 표시, 시간표 추정", "GET /api/v1/bus/arrivals, GET /api/v1/subway/arrivals", "TAGO + 서울시 지하철, 캐시 폴백"],
    ["6", "공통", "예외 처리", "ER_01_01", "검색 가능한 경로 없음", "Web/App", "후보 경로 없음", "운행 종료/경로 단절 구분 안내", "POST /api/v1/routes/recommend", "-"],
    ["7", "공통", "예외 처리", "ER_01_01", "네트워크 오류", "Web/App", "통신 실패 재시도", "마지막 조회 시각과 함께 기존 결과 유지", "공통 Envelope error", "-"],
    ["8", "홈", "메인 화면", "HM_01_01", "홈페이지", "Web/App", "출발지·목적지 입력", "양쪽 설정 시 경로 탐색 버튼 활성화", "POST /api/v1/routes/recommend", "-"],
    ["9", "홈", "출발지 설정", "HM_01_01", "출발지 변경", "Web/App", "현재위치/검색 선택", "출발지 변경 시 기존 경로 초기화", "GET /api/v1/places/search", "-"],
    ["10", "홈", "목적지 설정", "HM_01_01", "목적지 입력", "Web/App", "목적지 선택", "출발지와 동일하면 탐색 제한", "POST /api/v1/routes/recommend VALID_400", "-"],
    ["11", "홈", "검색 이력", "HM_01_01", "최근 검색", "Web/App", "최근 검색 장소 목록", "선택/삭제, 없으면 안내 문구", "GET/DELETE /api/v1/places/recent", "-"],
    ["12", "장소 검색", "검색", "SC_01_01", "장소 검색", "Web/App", "검색 입력창", "주소·장소명·역명·정류장명 검색", "GET /api/v1/places/search", "-"],
    ["13", "장소 검색", "결과", "SC_01_01", "검색 결과", "Web/App", "장소명·주소·좌표", "출발지/목적지 반영", "GET /api/v1/places/search", "-"],
    ["14", "장소 검색", "오류", "SC_01_01", "결과 없음", "Web/App", "0건/API 실패", "검색어 유지, 재시도", "GET /api/v1/places/search", "-"],
    ["15", "장소 검색", "교체", "HM_01_01", "출발지·목적지 교체", "Web/App", "상호 변경", "경로 결과 초기화", "클라이언트", "-"],
    ["16", "경로 탐색", "결과 조회", "RT_01_01", "경로 탐색", "Web/App", "탐색 프로세스", "검증>후보>실시간>도보/대기/환승>ETA>정렬", "POST /api/v1/routes/recommend", "버스+지하철 반영"],
    ["17", "경로 탐색", "후보 생성", "RT_01_01", "후보 경로", "Web/App", "도보/버스/지하철/복합", "운행 불가 후보 제외", "Router + GraphProvider", "버스 구간 TAGO, 지하철 구간 서울시"],
    ["18", "경로 탐색", "실시간 버스", "RT_01_01", "버스 도착예정", "Web/App", "노선번호·정류장ID 조회", "arrtime(초), arrprevstationcnt, vehicletp", "GET /api/v1/bus/arrivals, /arrivals/route", "TAGO getSttnAcctoArvlPrearngeInfoList / getSttnAcctoSpcifyRouteBusArvlPrearngeInfoList"],
    ["19", "경로 탐색", "실시간 지하철", "RT_01_01", "열차 도착예정", "Web/App", "노선·역·방면 조회", "barvlDt(초), arvlCd, lstcarAt, btrainSttus", "GET /api/v1/subway/arrivals", "서울시 realtimeStationArrival/ALL"],
    ["20", "경로 탐색", "도보 시간", "RT_01_01", "도보 소요시간", "Web/App", "걷기 속도 적용", "slow/normal/fast", "PATCH /api/v1/settings walk_speed", "-"],
    ["21", "경로 탐색", "신호등", "RT_01_01", "횡단 대기(선택)", "Web/App", "추후 고도화", "데이터 없으면 미반영", "미구현(명세 선택)", "-"],
    ["22", "경로 탐색", "승차 가능", "RT_01_01", "탑승 가능 차량", "Web/App", "여유시간 비교", "BOARDING_BUFFER_SECONDS", "RouteService._apply_realtime", "버스 arrtime / 지하철 barvlDt"],
    ["23", "경로 탐색", "환승 도보", "RT_01_01", "환승 이동시간", "Web/App", "하차 후 다음 승차 도보", "걷기 속도 적용", "Router walk step", "-"],
    ["24", "경로 탐색", "환승 대기", "RT_01_01", "다음 차량 대기", "Web/App", "후속 도착 - 환승지 도착", "실시간 도착정보 재조회", "버스/지하철 arrivals", "TAGO + 서울시"],
    ["25", "경로 탐색", "환승 연결", "RT_01_01", "환승 성공/실패", "Web/App", "여유시간 검증", "불가 시 다음 차량", "RouteService", "-"],
    ["26", "경로 탐색", "차량 이동", "RT_01_01", "구간 소요시간", "Web/App", "실시간/시간표 구분", "ETA heuristic + 실시간 대기", "POST /api/v1/routes/recommend", "-"],
    ["27", "경로 탐색", "최종 ETA", "RT_01_01", "총 소요·도착시각", "Web/App", "초 단위 합산", "도보+대기+이동+환승", "RouteSummary.eta", "-"],
    ["28", "경로 탐색", "최적 추천", "RT_01_01", "가장 빠른 경로", "Web/App", "도착시각 정렬", "동률 시 환승/도보", "HeuristicPredictor.rank_routes", "-"],
    ["29", "경로 탐색", "비교 카드", "RT_01_01", "경로 카드", "Web/App", "소요·환승·도보·노선", "실시간/시간표 표기", "RouteRecommendResponse", "-"],
    ["30", "경로 탐색", "갱신", "RT_01_01", "교통정보 갱신", "Web/App", "새로고침", "마지막 조회 시각 표시", "POST /api/v1/routes/{id}/refresh", "버스+지하철 재조회"],
    ["31", "경로 상세", "지도", "RT_02_01", "경로 상세", "Web/App", "단계별 지도", "도보/버스/지하철 구분", "GET /api/v1/routes/{id}", "-"],
    ["32", "경로 상세", "버스 탑승", "RT_02_01", "버스 상세", "Web/App", "노선·정류장·시간", "routeno, nodeid, arrtime, vehicletp", "transit_info + bus arrivals", "TAGO"],
    ["33", "경로 상세", "지하철 탑승", "RT_02_01", "지하철 상세", "Web/App", "노선·방면·역·시간", "trainLineNm, barvlDt, 급행", "transit_info + subway arrivals", "서울시"],
    ["34", "경로 상세", "환승 상세", "RT_02_01", "환승 안내", "Web/App", "위치·도보·대기", "환승 위험 경고", "GET /api/v1/routes/{id}", "-"],
    ["35", "경로 상세", "시간 요약", "RT_02_01", "구간별 시간", "Web/App", "총 소요=단계 합", "실시간 조회시각 표시", "RouteSummary", "-"],
    ["36", "경로 상세", "길안내 시작", "NV_01_01", "안내 시작", "App", "GPS 추적 시작", "권한 없으면 시작 제한", "클라이언트 + settings", "-"],
    ["37", "실시간 길안내", "이동 중", "NV_01_01", "현재 위치·단계", "App", "진행 단계 표시", "이탈 시 재탐색 안내", "POST /api/v1/routes/{id}/reroute", "-"],
    ["38", "실시간 길안내", "남은 시간", "NV_01_01", "남은 소요·ETA", "App", "위치+도착정보 재계산", "실패 시 직전값+시각", "refresh + arrivals", "버스+지하철"],
    ["39", "실시간 길안내", "재탐색", "NV_01_01", "대체 경로", "App", "사용자 동의 후 전환", "임의 전환 금지", "POST /api/v1/routes/{id}/reroute", "-"],
    ["40", "실시간 길안내", "환승 실패", "NV_01_01", "환승 경고", "App", "후속 차량 비교", "추가 소요 안내", "reroute + arrivals", "버스+지하철"],
    ["41", "실시간 길안내", "하차 안내", "NV_01_01", "하차·환승 알림", "App", "거리·시간 기준", "alert_minutes_before", "PATCH /api/v1/settings", "-"],
    ["42", "실시간 길안내", "종료", "NV_01_01", "안내 종료", "App", "추적·알림 중지", "목적지 도착 시에도 안내", "클라이언트", "-"],
    ["43", "설정", "걷기 속도", "ST_01_01", "이동 설정", "Web/App", "느림/보통/빠름", "이후 탐색부터 반영", "PATCH /api/v1/settings", "-"],
    ["44", "설정", "선호 조건", "ST_01_01", "경로 선호", "Web/App", "빠른도착/최소환승/최소도보", "route_sort", "PATCH /api/v1/settings", "-"],
    ["45", "설정", "알림", "ST_01_01", "환승·하차 알림", "App", "토글 + N분전", "alert_alight/transfer/minutes_before", "PATCH /api/v1/settings", "-"],
]
add_rows(ws, user_rows)
ws.column_dimensions["H"].width = 40
ws.column_dimensions["I"].width = 40
ws.column_dimensions["J"].width = 50

# --- 버스 ---
wsb = wb.create_sheet("버스도착정보_TAGO")
style_header(wsb, ["Page ID", "기능", "외부 API", "요청", "응답 필드", "route-app API"])
bus_rows = [
    ["BA_00_01", "버스도착정보조회 서비스 개요", "http://apis.data.go.kr/1613000/ArvlInfoInqireService", "serviceKey, REST GET, XML/JSON, 10~20초 갱신", "-", "-"],
    ["BA_01_01", "정류소별 도착예정정보 목록 조회", "GET .../getSttnAcctoArvlPrearngeInfoList", "serviceKey, cityCode, nodeId, numOfRows, pageNo, _type", "nodeid, nodenm, routeid, routeno, routetp, arrprevstationcnt, vehicletp, arrtime(초)", "GET /api/v1/bus/arrivals"],
    ["BA_02_01", "정류소별 특정노선버스 도착예정정보", "GET .../getSttnAcctoSpcifyRouteBusArvlPrearngeInfoList", "위 + routeId", "동일 응답 필드", "GET /api/v1/bus/arrivals/route"],
    ["BA_03_01", "도시코드 목록 조회", "GET .../getCtyCodeList", "serviceKey, _type", "citycode, cityname", "GET /api/v1/bus/cities"],
    ["BA_09_01", "OpenAPI 에러 코드", "XML OpenAPI_ServiceResponse", "portal 1,4,12,20,22,30,31,32,99 / provider 99 INVALID_REQUEST_PARAMETER_ERROR", "bus_api_error_codes", "Envelope error EXTERNAL_502 + DB 캐시 폴백"],
]
add_rows(wsb, bus_rows)
for col in "ABCDEF":
    wsb.column_dimensions[col].width = 28

# --- 지하철 ---
wss = wb.create_sheet("지하철실시간도착_서울시")
style_header(wss, ["Page ID", "기능", "외부 API", "요청", "응답 필드", "route-app API"])
sub_rows = [
    ["SM_00_01", "지하철 실시간 도착정보(일괄) 개요", "서울시 TOPIS", "KEY, TYPE, SERVICE, START_INDEX, END_INDEX", "전체역 실시간 도착", "-"],
    ["SM_01_01", "일괄 조회", "GET {base}/{KEY}/{TYPE}/realtimeStationArrival/ALL/{START}/{END}", "최대 1000건, sample키 1~5건", "subwayId, updnLine, trainLineNm, statnFid, statnTid, statnId, statnNm, trnsitCo, ordkey, subwayList, statnList, btrainSttus, barvlDt, btrainNo, bstatnId, bstatnNm, recptnDt, arvlMsg2, arvlMsg3, arvlCd, lstcarAt", "GET /api/v1/subway/arrivals"],
    ["SM_02_01", "에러/정보 메시지", "INFO-000, INFO-100, INFO-200, ERROR-300~336, ERROR-500/600/601", "페이지 분할 필요 시 ERROR-336", "subway_api_error_codes", "INFO-200=빈목록, 그 외 ExternalServiceError, 캐시 폴백"],
    ["SM_09_01", "원본 미기재 항목", "Call Back URL 등은 샘플 URL로 대체", "http://swopenAPI.seoul.go.kr/api/subway/{KEY}/json/realtimeStationArrival/ALL/{start}/{end}", "순번4 없음(3→5)", "SEOUL_SUBWAY_BASE_URL 설정"],
]
add_rows(wss, sub_rows)
for col in "ABCDEF":
    wss.column_dimensions[col].width = 32

# --- API 목록 ---
wsa = wb.create_sheet("백엔드API")
style_header(wsa, ["Method", "Path", "설명", "명세 ID"])
api_rows = [
    ["POST", "/api/v1/auth/signup", "회원가입 + 기본 이동설정", "ST_01_01"],
    ["POST", "/api/v1/auth/login", "로그인", "-"],
    ["POST", "/api/v1/auth/refresh", "토큰 갱신", "-"],
    ["GET", "/api/v1/places/search", "장소 검색(역/정류장 포함)", "SC_01_01"],
    ["GET", "/api/v1/places/recent", "최근 검색", "HM_01_01"],
    ["DELETE", "/api/v1/places/recent/{id}", "최근 검색 삭제", "HM_01_01"],
    ["POST", "/api/v1/routes/recommend", "최적 경로 추천(버스+지하철 실시간 반영)", "RT_01_01"],
    ["GET", "/api/v1/routes/{id}", "경로 상세", "RT_02_01"],
    ["POST", "/api/v1/routes/{id}/refresh", "실시간 교통정보 재조회", "RT_01_01"],
    ["POST", "/api/v1/routes/{id}/reroute", "이동 중 재탐색", "NV_01_01"],
    ["POST", "/api/v1/routes/{id}/feedback", "경로 피드백", "-"],
    ["GET", "/api/v1/bus/cities", "TAGO 도시코드", "BA_03_01"],
    ["GET", "/api/v1/bus/arrivals", "정류소별 버스 도착", "BA_01_01"],
    ["GET", "/api/v1/bus/arrivals/route", "특정노선 버스 도착", "BA_02_01"],
    ["GET", "/api/v1/subway/lines", "지하철 노선 마스터(17개)", "SM_01_01"],
    ["GET", "/api/v1/subway/arrivals", "지하철 실시간 도착(일괄 후 필터)", "SM_01_01"],
    ["POST", "/api/v1/subway/sync", "전체역 도착정보 적재", "SM_01_01"],
    ["GET", "/api/v1/settings", "걷기속도/선호/알림", "ST_01_01"],
    ["PATCH", "/api/v1/settings", "이동 설정 변경", "ST_01_01"],
    ["GET", "/health", "헬스체크", "-"],
]
add_rows(wsa, api_rows)
wsa.column_dimensions["B"].width = 42
wsa.column_dimensions["C"].width = 44

# --- DB ---
wsd = wb.create_sheet("PostgreSQL")
style_header(wsd, ["테이블", "출처", "핵심 컬럼"])
db_rows = [
    ["users / user_preferences", "기능정의서 설정", "walk_speed, route_sort, alert_alight, alert_transfer, alert_minutes_before"],
    ["places / recent_searches / favorites", "장소검색", "name, address, lat, lng, location(geography)"],
    ["routes / route_steps / step_segments", "경로 탐색", "transit_info JSONB에 버스/지하철 실시간 스냅샷"],
    ["bus_city_codes", "TAGO getCtyCodeList", "city_code, city_name"],
    ["bus_realtime_arrivals", "TAGO 도착정보", "node_id, route_id, routeno, arr_prev_station_cnt, vehicle_type, arr_time_seconds"],
    ["bus_api_error_codes", "TAGO 에러", "portal/provider 코드"],
    ["subway_lines", "서울시 subwayId", "1001~1032 17개 시드"],
    ["subway_stations / subway_station_transfers", "statnId, subwayList/statnList", "환승 매핑"],
    ["subway_realtime_arrivals", "realtimeStationArrival/ALL", "barvlDt, arvlCd, lstcarAt, ordkey 파싱 컬럼"],
    ["subway_api_error_codes", "서울시 INFO/ERROR", "retryable"],
    ["raw_public_data / data_sync_logs", "원문 보관", "body, result_code, fetched_at"],
]
add_rows(wsd, db_rows)
wsd.column_dimensions["A"].width = 40
wsd.column_dimensions["C"].width = 70

wb.save(out)
print("wrote", out)
